from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from .decorators import system_admin_required, technician_required, normal_user_required
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.forms import PasswordChangeForm
from django.db.models import Count, ProtectedError, Sum, F, ExpressionWrapper, DecimalField, Q, Avg
from django.db.models.functions import Coalesce, TruncMonth
from django.db import transaction
from django.utils import timezone
from django.contrib import messages
import csv
from django.http import HttpResponse
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from .forms import UserRegistrationForm
from .models import Asset, Allocation, Maintenance, Category, Department, Profile, SystemSettings, DamagedEquipment, UserRequest, UserFeedback

@login_required
def dashboard(request):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    
    if role == 'NORMAL_USER':
        # NORMAL_USER Statistics
        my_equipment_count = request.user.allocations.filter(status__in=['Active', 'Partially Returned']).count()
        active_requests_count = UserRequest.objects.filter(user=request.user, status__in=['PENDING', 'ASSIGNED', 'IN_PROGRESS']).count()
        resolved_requests_count = UserRequest.objects.filter(user=request.user, status='RESOLVED').count()
        open_issues_count = DamagedEquipment.objects.filter(reported_by=request.user, status__in=['REPORTED', 'UNDER_REVIEW', 'UNDER_REPAIR']).count()
        
        context = {
            'my_equipment_count': my_equipment_count,
            'active_requests_count': active_requests_count,
            'resolved_requests_count': resolved_requests_count,
            'open_issues_count': open_issues_count
        }
        return render(request, 'assets/user_dashboard.html', context)
        
    elif role == 'TECHNICIAN':
        recent_maintenance = Maintenance.objects.select_related('asset').order_by('-maintenance_date', '-id')[:5]
        maintenance_assets = Asset.objects.filter(status='Under Maintenance').aggregate(m=Sum('quantity'))['m'] or 0
        total_assets = Asset.objects.aggregate(total=Sum('quantity'))['total'] or 0
        context = {
            'recent_maintenance': recent_maintenance,
            'maintenance_assets': maintenance_assets,
            'total_assets': total_assets,
            'damaged_count': DamagedEquipment.objects.filter(status__in=['REPORTED', 'UNDER_REVIEW', 'UNDER_REPAIR']).count(),
            'pending_requests': UserRequest.objects.filter(status='PENDING').count(),
            'in_progress_requests': UserRequest.objects.filter(status__in=['ASSIGNED', 'IN_PROGRESS']).count()
        }
        return render(request, 'assets/technician_dashboard.html', context)

    total_assets = Asset.objects.count()
    total_asset_quantity = Asset.objects.aggregate(total=Sum('quantity'))['total'] or 0
    available_quantity = Asset.objects.aggregate(avail=Sum('available_quantity'))['avail'] or 0
    allocated_quantity = total_asset_quantity - available_quantity
    maintenance_assets = Asset.objects.filter(status='Under Maintenance').count()
    damaged_equipment = DamagedEquipment.objects.filter(status__in=['REPORTED', 'UNDER_REVIEW', 'UNDER_REPAIR']).count()

    active_allocations = Allocation.objects.filter(status__in=['Active', 'Partially Returned']).count()
    pending_returns = Allocation.objects.filter(status__in=['Active', 'Partially Returned'], expected_return_date__isnull=False).count()
    pending_requests = UserRequest.objects.filter(status='PENDING').count()
    resolved_requests = UserRequest.objects.filter(status='RESOLVED').count()
    
    total_users = User.objects.filter(profile__role='NORMAL_USER').count()
    total_technicians = User.objects.filter(profile__role='TECHNICIAN').count()

    recent_allocations = Allocation.objects.select_related('asset', 'department').order_by('-allocation_date', '-id')[:5]
    recent_support_requests = UserRequest.objects.select_related('user', 'asset', 'assigned_technician').order_by('-created_at')[:5]
    recent_damage_reports = DamagedEquipment.objects.select_related('reported_by', 'asset', 'assigned_technician').order_by('-created_at')[:5]
    recent_maintenance = Maintenance.objects.select_related('asset').order_by('-maintenance_date', '-id')[:5]
    
    context = {
        'total_assets': total_assets,
        'total_asset_quantity': total_asset_quantity,
        'available_quantity': available_quantity,
        'allocated_quantity': allocated_quantity,
        'maintenance_assets': maintenance_assets,
        'damaged_equipment': damaged_equipment,
        
        'active_allocations': active_allocations,
        'pending_returns': pending_returns,
        'pending_requests': pending_requests,
        'open_damage_reports': damaged_equipment, # Same as damaged_equipment
        'resolved_requests': resolved_requests,
        
        'total_users': total_users,
        'total_technicians': total_technicians,
        
        'recent_allocations': recent_allocations,
        'recent_support_requests': recent_support_requests,
        'recent_damage_reports': recent_damage_reports,
        'recent_maintenance': recent_maintenance,
    }
    return render(request, 'assets/dashboard.html', context)

@login_required
@technician_required
def asset_list(request):
    assets = Asset.objects.select_related('category').order_by('-created_at')
    categories = Category.objects.all().order_by('name')
    
    total_assets = assets.aggregate(total=Sum('quantity'))['total'] or 0
    available = assets.aggregate(avail=Sum('available_quantity'))['avail'] or 0
    allocated = total_assets - available
    under_maintenance = assets.filter(status='Under Maintenance').aggregate(m=Sum('quantity'))['m'] or 0
    
    context = {
        'assets': assets,
        'categories': categories,
        'condition_choices': Asset.CONDITION_CHOICES,
        'status_choices': Asset.STATUS_CHOICES,
        'total_assets': total_assets,
            'damaged_count': DamagedEquipment.objects.filter(status__in=['REPORTED', 'UNDER_REVIEW', 'UNDER_REPAIR']).count(),
            'pending_requests': UserRequest.objects.filter(status='PENDING').count(),
            'in_progress_requests': UserRequest.objects.filter(status__in=['ASSIGNED', 'IN_PROGRESS']).count(),
        'available_count': available,
        'allocated_count': allocated,
        'maintenance_count': under_maintenance,
    }
    return render(request, 'assets/asset_list.html', context)

@login_required
@system_admin_required
def asset_add(request):
    if request.method == 'POST':
        asset_code = request.POST.get('asset_code')
        name = request.POST.get('name')
        category_id = request.POST.get('category_id')
        description = request.POST.get('description', '')
        purchase_date = request.POST.get('purchase_date') or None
        purchase_cost = request.POST.get('purchase_cost') or None
        quantity = request.POST.get('quantity') or 1
        available_quantity = request.POST.get('available_quantity') or 1
        location = request.POST.get('location', '')
        condition = request.POST.get('condition', 'Good')
        status = request.POST.get('status', 'Available')
        
        if asset_code and name and category_id:
            if Asset.objects.filter(asset_code=asset_code).exists():
                messages.error(request, f"Asset code '{asset_code}' already exists.")
            else:
                category = get_object_or_404(Category, pk=category_id)
                Asset.objects.create(
                    asset_code=asset_code,
                    name=name,
                    category=category,
                    description=description,
                    purchase_date=purchase_date,
                    purchase_cost=purchase_cost,
                    quantity=quantity,
                    available_quantity=available_quantity,
                    location=location,
                    condition=condition,
                    status=status
                )
                messages.success(request, f"Asset '{name}' added successfully.")
    return redirect('assets:asset_list')

@login_required
@system_admin_required
def asset_edit(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=pk)
            asset_code = request.POST.get('asset_code')
            name = request.POST.get('name')
            category_id = request.POST.get('category_id')
            description = request.POST.get('description', '')
            purchase_date = request.POST.get('purchase_date') or None
            purchase_cost = request.POST.get('purchase_cost') or None
            new_quantity = int(request.POST.get('quantity') or 1)
            location = request.POST.get('location', '')
            condition = request.POST.get('condition', 'Good')
            status = request.POST.get('status', 'Available')
            
            if asset_code and name and category_id:
                if Asset.objects.filter(asset_code=asset_code).exclude(pk=pk).exists():
                    messages.error(request, f"Asset code '{asset_code}' already exists.")
                else:
                    allocated_quantity = asset.quantity - asset.available_quantity
                    
                    if new_quantity < allocated_quantity:
                        messages.error(request, f"Cannot reduce total quantity to {new_quantity}. {allocated_quantity} units are currently allocated.")
                        return redirect('assets:asset_list')
                        
                    asset.asset_code = asset_code
                    asset.name = name
                    asset.category = get_object_or_404(Category, pk=category_id)
                    asset.description = description
                    asset.purchase_date = purchase_date
                    asset.purchase_cost = purchase_cost
                    asset.quantity = new_quantity
                    asset.available_quantity = new_quantity - allocated_quantity
                    asset.location = location
                    asset.condition = condition
                    
                    # Temporarily save manually passed status, but recalculate properly
                    asset.status = status
                    asset.save()
                    
                    # Safely apply status based on maintenance and allocations
                    update_asset_status_for_maintenance(asset)
                    
                    messages.success(request, f"Asset '{name}' updated successfully.")
    return redirect('assets:asset_list')

@login_required
@system_admin_required
def asset_delete(request, pk):
    if request.method == 'POST':
        asset = get_object_or_404(Asset, pk=pk)
        try:
            name = asset.name
            asset.delete()
            messages.success(request, f"Asset '{name}' deleted successfully.")
        except ProtectedError:
            messages.error(request, f"Cannot delete asset '{asset.name}' because it has related allocations or maintenance records.")
    return redirect('assets:asset_list')

@login_required
@system_admin_required
def category_list(request):
    categories = Category.objects.annotate(asset_count=Count('asset')).order_by('-created_at')
    return render(request, 'assets/category_list.html', {'categories': categories})

@login_required
@system_admin_required
def category_add(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        description = request.POST.get('description', '')
        if name:
            if Category.objects.filter(name=name).exists():
                messages.error(request, f"Category '{name}' already exists.")
            else:
                Category.objects.create(name=name, description=description)
                messages.success(request, f"Category '{name}' added successfully.")
    return redirect('assets:category_list')

@login_required
@system_admin_required
def category_edit(request, pk):
    if request.method == 'POST':
        category = get_object_or_404(Category, pk=pk)
        name = request.POST.get('name')
        description = request.POST.get('description', '')
        if name:
            if Category.objects.filter(name=name).exclude(pk=pk).exists():
                messages.error(request, f"Category '{name}' already exists.")
            else:
                category.name = name
                category.description = description
                category.save()
                messages.success(request, f"Category '{name}' updated successfully.")
    return redirect('assets:category_list')

@login_required
@system_admin_required
def category_delete(request, pk):
    if request.method == 'POST':
        category = get_object_or_404(Category, pk=pk)
        try:
            name = category.name
            category.delete()
            messages.success(request, f"Category '{name}' deleted successfully.")
        except ProtectedError:
            messages.error(request, f"Cannot delete category '{category.name}' because it contains assets.")
    return redirect('assets:category_list')

@login_required
@system_admin_required
def department_list(request):
    departments = Department.objects.annotate(allocation_count=Count('allocation')).order_by('-created_at')
    return render(request, 'assets/department_list.html', {'departments': departments})

@login_required
@system_admin_required
def department_add(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        code = request.POST.get('code')
        location = request.POST.get('location', '')
        if name and code:
            if Department.objects.filter(name=name).exists():
                messages.error(request, f"Department '{name}' already exists.")
            elif Department.objects.filter(code=code).exists():
                messages.error(request, f"Department code '{code}' already exists.")
            else:
                Department.objects.create(name=name, code=code, location=location)
                messages.success(request, f"Department '{name}' added successfully.")
    return redirect('assets:department_list')

@login_required
@system_admin_required
def department_edit(request, pk):
    if request.method == 'POST':
        department = get_object_or_404(Department, pk=pk)
        name = request.POST.get('name')
        code = request.POST.get('code')
        location = request.POST.get('location', '')
        if name and code:
            if Department.objects.filter(name=name).exclude(pk=pk).exists():
                messages.error(request, f"Department '{name}' already exists.")
            elif Department.objects.filter(code=code).exclude(pk=pk).exists():
                messages.error(request, f"Department code '{code}' already exists.")
            else:
                department.name = name
                department.code = code
                department.location = location
                department.save()
                messages.success(request, f"Department '{name}' updated successfully.")
    return redirect('assets:department_list')

@login_required
@system_admin_required
def department_delete(request, pk):
    if request.method == 'POST':
        department = get_object_or_404(Department, pk=pk)
        try:
            name = department.name
            department.delete()
            messages.success(request, f"Department '{name}' deleted successfully.")
        except ProtectedError:
            messages.error(request, f"Cannot delete department '{department.name}' because it has active allocations.")
    return redirect('assets:department_list')

@login_required
@technician_required
def allocation_list(request):
    allocations = Allocation.objects.select_related('asset', 'department', 'assigned_user').order_by('-allocation_date', '-id')
    assets = Asset.objects.filter(available_quantity__gt=0).order_by('name')
    departments = Department.objects.all().order_by('name')
    users = User.objects.filter(is_active=True).order_by('first_name')
    
    total_allocations = allocations.count()
    active_allocations = allocations.filter(status='Active').count()
    returned_allocations = allocations.filter(status__in=['Partially Returned', 'Returned']).count()
    
    context = {
        'allocations': allocations,
        'assets': assets,
        'departments': departments,
        'users': users,
        'status_choices': Allocation.STATUS_CHOICES,
        'total_allocations': total_allocations,
        'active_allocations': active_allocations,
        'returned_allocations': returned_allocations,
    }
    return render(request, 'assets/allocation_list.html', context)

@login_required
@system_admin_required
def allocation_add(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset_id')
        department_id = request.POST.get('department_id')
        assigned_to = request.POST.get('assigned_to', '')
        assigned_user_id = request.POST.get('assigned_user_id', '')
        assigned_user = User.objects.get(pk=assigned_user_id) if assigned_user_id else None
        allocation_type = request.POST.get('allocation_type', 'PERMANENT')
        quantity = int(request.POST.get('quantity', 0))
        allocation_date = request.POST.get('allocation_date')
        expected_return_date = request.POST.get('expected_return_date') or None
        purpose = request.POST.get('purpose', '')
        
        if quantity <= 0:
            messages.error(request, "Quantity must be greater than zero.")
            return redirect('assets:allocation_list')
            
        with transaction.atomic():
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=asset_id)
            department = get_object_or_404(Department, pk=department_id)
            
            if quantity > asset.available_quantity:
                messages.error(request, f"Requested quantity ({quantity}) exceeds available quantity ({asset.available_quantity}).")
            else:
                Allocation.objects.create(
                    asset=asset,
                    department=department,
                    assigned_to=assigned_to,
                    assigned_user=assigned_user,
                    allocation_type=allocation_type,
                    quantity=quantity,
                    allocation_date=allocation_date,
                    expected_return_date=expected_return_date,
                    purpose=purpose,
                    status='Active'
                )
                
                asset.available_quantity -= quantity
                if asset.available_quantity == 0:
                    asset.status = 'Fully Allocated'
                elif asset.available_quantity < asset.quantity:
                    asset.status = 'Partially Allocated'
                asset.save()
                
                messages.success(request, f"Successfully allocated {quantity}x {asset.name} to {department.name}.")
    return redirect('assets:allocation_list')

@login_required
@system_admin_required
def allocation_edit(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            allocation = get_object_or_404(Allocation.objects.select_for_update(), pk=pk)
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=allocation.asset_id)
            
            department_id = request.POST.get('department_id')
            assigned_to = request.POST.get('assigned_to', '')
            assigned_user_id = request.POST.get('assigned_user_id', '')
            assigned_user = User.objects.get(pk=assigned_user_id) if assigned_user_id else None
            allocation_type = request.POST.get('allocation_type', allocation.allocation_type)
            new_quantity = int(request.POST.get('quantity', 0))
            allocation_date = request.POST.get('allocation_date')
            expected_return_date = request.POST.get('expected_return_date') or None
            purpose = request.POST.get('purpose', '')
            
            if new_quantity <= 0:
                messages.error(request, "Quantity must be greater than zero.")
                return redirect('assets:allocation_list')
                
            if new_quantity < allocation.returned_quantity:
                messages.error(request, f"Cannot decrease quantity below already returned quantity ({allocation.returned_quantity}).")
                return redirect('assets:allocation_list')
            
            if allocation.status in ['Active', 'Partially Returned']:
                qty_difference = new_quantity - allocation.quantity
                
                if qty_difference > asset.available_quantity:
                    messages.error(request, f"Cannot increase quantity by {qty_difference}. Only {asset.available_quantity} available.")
                    return redirect('assets:allocation_list')
                    
                asset.available_quantity -= qty_difference
                
                # Strict boundary enforcement
                if asset.available_quantity > asset.quantity:
                    asset.available_quantity = asset.quantity
                if asset.available_quantity < 0:
                    asset.available_quantity = 0
                    
                if asset.available_quantity == 0:
                    asset.status = 'Fully Allocated'
                elif asset.available_quantity == asset.quantity:
                    asset.status = 'Available'
                else:
                    asset.status = 'Partially Allocated'
                asset.save()
            
            allocation.department_id = department_id
            allocation.assigned_to = assigned_to
            allocation.allocation_type = allocation_type
            allocation.quantity = new_quantity
            allocation.allocation_date = allocation_date
            allocation.expected_return_date = expected_return_date
            allocation.purpose = purpose
            
            if allocation.returned_quantity == allocation.quantity:
                allocation.status = 'Returned'
            elif allocation.returned_quantity > 0:
                allocation.status = 'Partially Returned'
            else:
                allocation.status = 'Active'
                
            allocation.save()
            
            messages.success(request, f"Allocation updated successfully.")
    return redirect('assets:allocation_list')

@login_required
@system_admin_required
def allocation_return(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            allocation = get_object_or_404(Allocation.objects.select_for_update(), pk=pk)
            if allocation.status == 'Returned':
                messages.error(request, "This allocation is already fully returned.")
                return redirect('assets:allocation_list')
                
            actual_return_date = request.POST.get('actual_return_date') or timezone.now().date()
            active_qty = allocation.quantity - allocation.returned_quantity
            return_quantity = int(request.POST.get('return_quantity', active_qty))
            
            if return_quantity <= 0:
                messages.error(request, "Return quantity must be greater than zero.")
                return redirect('assets:allocation_list')
            
            if return_quantity > active_qty:
                messages.error(request, f"Cannot return {return_quantity}. Only {active_qty} are currently active.")
                return redirect('assets:allocation_list')
            
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=allocation.asset_id)
            asset.available_quantity += return_quantity
            if asset.available_quantity > asset.quantity:
                asset.available_quantity = asset.quantity
                
            if asset.available_quantity == asset.quantity:
                asset.status = 'Available'
            else:
                asset.status = 'Partially Allocated'
            asset.save()
            
            allocation.returned_quantity += return_quantity
            allocation.actual_return_date = actual_return_date
            
            if allocation.returned_quantity == allocation.quantity:
                allocation.status = 'Returned'
            else:
                allocation.status = 'Partially Returned'
                
            allocation.save()
            
            messages.success(request, f"Successfully returned {return_quantity} units.")
    return redirect('assets:allocation_list')

@login_required
@system_admin_required
def allocation_delete(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            allocation = get_object_or_404(Allocation.objects.select_for_update(), pk=pk)
            
            if allocation.status != 'Returned':
                asset = get_object_or_404(Asset.objects.select_for_update(), pk=allocation.asset_id)
                active_qty = allocation.quantity - allocation.returned_quantity
                asset.available_quantity += active_qty
                if asset.available_quantity > asset.quantity:
                    asset.available_quantity = asset.quantity
                    
                if asset.available_quantity == asset.quantity:
                    asset.status = 'Available'
                else:
                    asset.status = 'Partially Allocated'
                asset.save()
                
            allocation.delete()
            messages.success(request, "Allocation deleted successfully.")
    return redirect('assets:allocation_list')

def update_asset_status_for_maintenance(asset):
    has_active_maintenance = Maintenance.objects.filter(
        asset=asset, 
        status__in=['Pending', 'In Progress']
    ).exists()
    
    if has_active_maintenance:
        asset.status = 'Under Maintenance'
    else:
        if asset.available_quantity == asset.quantity:
            asset.status = 'Available'
        elif asset.available_quantity == 0:
            asset.status = 'Fully Allocated'
        else:
            asset.status = 'Partially Allocated'
    asset.save()

@login_required
@technician_required
def maintenance_list(request):
    maintenances = Maintenance.objects.select_related('asset').order_by('-maintenance_date', '-id')
    assets = Asset.objects.all().order_by('name')
    
    total_records = maintenances.count()
    in_progress = maintenances.filter(status='In Progress').count()
    completed = maintenances.filter(status='Completed').count()
    total_cost = maintenances.aggregate(total=Sum('cost'))['total'] or 0
    
    context = {
        'maintenances': maintenances,
        'assets': assets,
        'status_choices': Maintenance.STATUS_CHOICES,
        'total_records': total_records,
        'in_progress': in_progress,
        'completed': completed,
        'total_cost': total_cost,
    }
    return render(request, 'assets/maintenance_list.html', context)

@login_required
@system_admin_required
def maintenance_add(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset_id')
        description = request.POST.get('description', '')
        maintenance_date = request.POST.get('maintenance_date')
        cost = request.POST.get('cost') or 0
        technician = request.POST.get('technician', '')
        status = request.POST.get('status', 'Pending')
        
        with transaction.atomic():
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=asset_id)
            
            Maintenance.objects.create(
                asset=asset,
                description=description,
                maintenance_date=maintenance_date,
                cost=cost,
                technician=technician,
                status=status
            )
            
            update_asset_status_for_maintenance(asset)
            messages.success(request, f"Maintenance record for {asset.name} added successfully.")
            
    return redirect('assets:maintenance_list')

@login_required
@system_admin_required
def maintenance_edit(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            maintenance = get_object_or_404(Maintenance.objects.select_for_update(), pk=pk)
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=maintenance.asset_id)
            
            description = request.POST.get('description', '')
            maintenance_date = request.POST.get('maintenance_date')
            cost = request.POST.get('cost') or 0
            technician = request.POST.get('technician', '')
            status = request.POST.get('status', 'Pending')
            
            maintenance.description = description
            maintenance.maintenance_date = maintenance_date
            maintenance.cost = cost
            maintenance.technician = technician
            maintenance.status = status
            maintenance.save()
            
            update_asset_status_for_maintenance(asset)
            messages.success(request, "Maintenance record updated successfully.")
            
    return redirect('assets:maintenance_list')

@login_required
@system_admin_required
def maintenance_delete(request, pk):
    if request.method == 'POST':
        with transaction.atomic():
            maintenance = get_object_or_404(Maintenance.objects.select_for_update(), pk=pk)
            asset = get_object_or_404(Asset.objects.select_for_update(), pk=maintenance.asset_id)
            
            maintenance.delete()
            update_asset_status_for_maintenance(asset)
            
            messages.success(request, "Maintenance record deleted successfully.")
            
    return redirect('assets:maintenance_list')

@login_required
def profile(request):
    profile_obj, created = Profile.objects.get_or_create(user=request.user)
    
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        email = request.POST.get('email', '')
        
        request.user.first_name = first_name
        request.user.last_name = last_name
        request.user.email = email
        request.user.save()
        
        if 'profile_picture' in request.FILES:
            profile_obj.profile_picture = request.FILES['profile_picture']
            profile_obj.save()
            
        messages.success(request, "Profile updated successfully.")
        return redirect('assets:profile')
        
    return render(request, 'assets/profile.html', {'profile': profile_obj})

@login_required
@system_admin_required
def settings_view(request):
    profile_obj, _ = Profile.objects.get_or_create(user=request.user)
    sys_settings, _ = SystemSettings.objects.get_or_create(id=1)
    
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'update_notifications':
            profile_obj.low_stock_notifications = request.POST.get('low_stock_notifications') == 'on'
            profile_obj.overdue_allocation_notifications = request.POST.get('overdue_allocation_notifications') == 'on'
            profile_obj.maintenance_notifications = request.POST.get('maintenance_notifications') == 'on'
            profile_obj.save()
            messages.success(request, "Notification preferences updated successfully.")
            return redirect('assets:settings')
            
        elif action == 'update_system' and (request.user.is_staff or request.user.is_superuser):
            threshold = request.POST.get('low_stock_threshold')
            try:
                val = int(threshold)
                if val >= 0:
                    sys_settings.low_stock_threshold = val
                    sys_settings.save()
                    messages.success(request, "System preferences updated successfully.")
                else:
                    messages.error(request, "Threshold must be a non-negative integer.")
            except (ValueError, TypeError):
                messages.error(request, "Invalid threshold value.")
            return redirect('assets:settings')

    context = {
        'profile': profile_obj,
        'sys_settings': sys_settings
    }
    return render(request, 'assets/settings.html', context)


def get_filtered_assets(request):
    f_category = request.GET.get('category')
    f_asset_status = request.GET.get('asset_status')
    
    asset_report = Asset.objects.select_related('category').annotate(
        total_value=ExpressionWrapper(
            Coalesce(F('purchase_cost'), 0.0, output_field=DecimalField()) * F('quantity'),
            output_field=DecimalField()
        ),
        allocated_qty=F('quantity') - F('available_quantity')
    ).order_by('name')

    if f_category:
        asset_report = asset_report.filter(category_id=f_category)
    if f_asset_status:
        asset_report = asset_report.filter(status=f_asset_status)
    return asset_report

def get_filtered_allocations(request):
    f_department = request.GET.get('department')
    f_allocation_status = request.GET.get('allocation_status')
    f_start_date = request.GET.get('start_date')
    f_end_date = request.GET.get('end_date')

    allocation_report = Allocation.objects.select_related('asset', 'department').order_by('-allocation_date')
    if f_department:
        allocation_report = allocation_report.filter(department_id=f_department)
    if f_allocation_status:
        allocation_report = allocation_report.filter(status=f_allocation_status)
    if f_start_date:
        allocation_report = allocation_report.filter(allocation_date__gte=f_start_date)
    if f_end_date:
        allocation_report = allocation_report.filter(allocation_date__lte=f_end_date)
    return allocation_report

def get_filtered_maintenance(request):
    f_maintenance_status = request.GET.get('maintenance_status')
    f_start_date = request.GET.get('start_date')
    f_end_date = request.GET.get('end_date')

    maintenance_report = Maintenance.objects.select_related('asset').order_by('-maintenance_date')
    if f_maintenance_status:
        maintenance_report = maintenance_report.filter(status=f_maintenance_status)
    if f_start_date:
        maintenance_report = maintenance_report.filter(maintenance_date__gte=f_start_date)
    if f_end_date:
        maintenance_report = maintenance_report.filter(maintenance_date__lte=f_end_date)
    return maintenance_report


@login_required
@system_admin_required
def reports(request):
    # 1. Asset Metrics
    total_assets = Asset.objects.count()
    
    asset_aggs = Asset.objects.aggregate(
        total_val=Sum(ExpressionWrapper(Coalesce(F('purchase_cost'), 0.0, output_field=DecimalField()) * F('quantity'), output_field=DecimalField())),
        total_qty=Sum('quantity'),
        avail_qty=Sum('available_quantity')
    )
    
    total_asset_value = asset_aggs['total_val'] or 0
    total_quantity = asset_aggs['total_qty'] or 0
    available_quantity = asset_aggs['avail_qty'] or 0
    allocated_quantity = total_quantity - available_quantity
    
    # 2. Allocation Metrics
    active_allocations = Allocation.objects.filter(status='Active').count()
    returned_allocations = Allocation.objects.filter(status='Returned').count()
    
    today = timezone.now().date()
    overdue_allocations = Allocation.objects.filter(
        expected_return_date__lt=today
    ).exclude(status='Returned').count()
    
    # 3. Maintenance Metrics
    total_maintenance_records = Maintenance.objects.count()
    active_maintenance = Maintenance.objects.filter(status__in=['Pending', 'In Progress']).count()
    completed_maintenance = Maintenance.objects.filter(status='Completed').count()
    total_maintenance_cost = Maintenance.objects.aggregate(total=Sum('cost'))['total'] or 0
    
    # --- CHARTS DATA ---
    
    # A. Assets by Category
    category_data = Category.objects.annotate(count=Count('asset'))
    category_labels = [c.name for c in category_data]
    category_counts = [c.count for c in category_data]
    
    # B. Assets by Status
    status_counts_db = Asset.objects.values('status').annotate(count=Count('id'))
    status_map = {item['status']: item['count'] for item in status_counts_db}
    asset_status_labels = [choice[1] for choice in Asset.STATUS_CHOICES]
    asset_status_counts = [status_map.get(choice[0], 0) for choice in Asset.STATUS_CHOICES]
    
    # C. Assets by Condition
    condition_counts_db = Asset.objects.values('condition').annotate(count=Count('id'))
    condition_map = {item['condition']: item['count'] for item in condition_counts_db}
    condition_labels = [choice[1] for choice in Asset.CONDITION_CHOICES]
    condition_counts = [condition_map.get(choice[0], 0) for choice in Asset.CONDITION_CHOICES]
    
    # D. Department-wise Allocations
    dept_data = Department.objects.annotate(
        allocated=Sum('allocation__quantity', filter=~Q(allocation__status='Returned'))
    )
    department_labels = [d.name for d in dept_data]
    department_allocation_quantities = [d.allocated or 0 for d in dept_data]
    
    # E. Maintenance Cost Analysis
    maint_data = Maintenance.objects.annotate(month=TruncMonth('maintenance_date')).values('month').annotate(total=Sum('cost')).order_by('month')
    maintenance_month_labels = [item['month'].strftime('%b %Y') for item in maint_data if item['month']]
    maintenance_month_costs = [float(item['total']) for item in maint_data]
    
    context = {
        'total_assets': total_assets,
            'damaged_count': DamagedEquipment.objects.filter(status__in=['REPORTED', 'UNDER_REVIEW', 'UNDER_REPAIR']).count(),
            'pending_requests': UserRequest.objects.filter(status='PENDING').count(),
            'in_progress_requests': UserRequest.objects.filter(status__in=['ASSIGNED', 'IN_PROGRESS']).count(),
        'total_asset_value': total_asset_value,
        'total_quantity': total_quantity,
        'available_quantity': available_quantity,
        'allocated_quantity': allocated_quantity,
        
        'active_allocations': active_allocations,
        'returned_allocations': returned_allocations,
        'overdue_allocations': overdue_allocations,
        
        'total_maintenance_records': total_maintenance_records,
        'active_maintenance': active_maintenance,
        'completed_maintenance': completed_maintenance,
        'total_maintenance_cost': total_maintenance_cost,
        
        # Charts Data
        'category_labels': category_labels,
        'category_counts': category_counts,
        
        'asset_status_labels': asset_status_labels,
        'asset_status_counts': asset_status_counts,
        
        'condition_labels': condition_labels,
        'condition_counts': condition_counts,
        
        'department_labels': department_labels,
        'department_allocation_quantities': department_allocation_quantities,
        
        'maintenance_month_labels': maintenance_month_labels,
        'maintenance_month_costs': maintenance_month_costs,
    }
    
    # --- DETAILED REPORTS (FILTERED) ---
    asset_report = get_filtered_assets(request)
    allocation_report = get_filtered_allocations(request)
    maintenance_report = get_filtered_maintenance(request)

        
    context.update({
        'asset_report': asset_report,
        'allocation_report': allocation_report,
        'maintenance_report': maintenance_report,
        
        # Filter options
        'categories': Category.objects.all(),
        'users': User.objects.filter(is_active=True).exclude(username='admin'),
        'departments': Department.objects.all(),
        'asset_status_choices': Asset.STATUS_CHOICES,
        'allocation_status_choices': Allocation.STATUS_CHOICES,
        'maintenance_status_choices': Maintenance.STATUS_CHOICES,
    })
    
    return render(request, 'assets/reports.html', context)

@login_required
def change_password(request):
    if request.method == 'POST':
        form = PasswordChangeForm(request.user, request.POST)
        if form.is_valid():
            user = form.save()
            update_session_auth_hash(request, user)  # Important to keep user logged in
            messages.success(request, 'Your password was successfully updated!')
            return redirect('assets:settings')
        else:
            messages.error(request, 'Please correct the error below.')
    else:
        form = PasswordChangeForm(request.user)
    return render(request, 'assets/change_password.html', {'form': form})


# ==============================================================================
# EXPORT VIEWS
# ==============================================================================

@login_required
@system_admin_required
def export_assets_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="asset_inventory_report_{datetime.now().strftime("%Y%m%d")}.csv"'
    response.write('\ufeff'.encode('utf8')) # utf-8-sig
    
    writer = csv.writer(response)
    writer.writerow(['Asset Code', 'Asset Name', 'Category', 'Total Quantity', 'Available Quantity', 'Allocated Quantity', 'Condition', 'Status', 'Purchase Cost', 'Total Value'])
    
    for obj in get_filtered_assets(request):
        writer.writerow([
            obj.asset_code, obj.name, obj.category.name, obj.quantity, obj.available_quantity, obj.allocated_qty,
            obj.get_condition_display(), obj.get_status_display(), obj.purchase_cost or '', obj.total_value or ''
        ])
    return response

@login_required
@system_admin_required
def export_assets_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asset Inventory"
    
    # Headers
    ws.append(['Campus Assets - Asset Inventory Report'])
    ws.append([f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'])
    ws.append([])
    
    headers = ['Asset Code', 'Asset Name', 'Category', 'Total Quantity', 'Available Quantity', 'Allocated Quantity', 'Condition', 'Status', 'Purchase Cost', 'Total Value']
    ws.append(headers)
    
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=4, column=col).font = header_font
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A4:J4"
    
    for obj in get_filtered_assets(request):
        row = [
            obj.asset_code, obj.name, obj.category.name, obj.quantity, obj.available_quantity, obj.allocated_qty,
            obj.get_condition_display(), obj.get_status_display(), obj.purchase_cost or 0.0, obj.total_value or 0.0
        ]
        ws.append(row)
        ws.cell(row=ws.max_row, column=9).number_format = '[$₹-en-IN]#,##0.00'
        ws.cell(row=ws.max_row, column=10).number_format = '[$₹-en-IN]#,##0.00'
        
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="asset_inventory_report_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response

@login_required
@system_admin_required
def export_assets_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="asset_inventory_report_{datetime.now().strftime("%Y%m%d")}.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Campus Assets", styles['Heading1']))
    elements.append(Paragraph("Asset Inventory Report", styles['Heading2']))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    data = [['Code', 'Name', 'Category', 'Qty', 'Avail', 'Alloc', 'Condition', 'Status', 'Unit Cost', 'Total Value']]
    for obj in get_filtered_assets(request):
        pc = f"Rs. {obj.purchase_cost:.2f}" if obj.purchase_cost else "-"
        tv = f"Rs. {obj.total_value:.2f}" if obj.total_value else "-"
        data.append([
            obj.asset_code, obj.name, obj.category.name, str(obj.quantity), str(obj.available_quantity), 
            str(obj.allocated_qty), obj.get_condition_display(), obj.get_status_display(), pc, tv
        ])
        
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A3626')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(t)
    doc.build(elements)
    return response

# -- Allocations Exports --
@login_required
@system_admin_required
def export_allocations_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="allocations_report_{datetime.now().strftime("%Y%m%d")}.csv"'
    response.write('\ufeff'.encode('utf8'))
    
    writer = csv.writer(response)
    writer.writerow(['Asset', 'Asset Code', 'Department', 'Assigned To', 'Quantity', 'Allocation Date', 'Expected Return Date', 'Actual Return Date', 'Purpose', 'Status'])
    
    for obj in get_filtered_allocations(request):
        writer.writerow([
            obj.asset.name, obj.asset.asset_code, obj.department.name, obj.assigned_to, obj.quantity,
            obj.allocation_date, obj.expected_return_date or '', obj.actual_return_date or '', obj.purpose, obj.get_status_display()
        ])
    return response

@login_required
@system_admin_required
def export_allocations_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Allocations"
    
    ws.append(['Campus Assets - Allocations Report'])
    ws.append([f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'])
    ws.append([])
    
    headers = ['Asset', 'Asset Code', 'Department', 'Assigned To', 'Quantity', 'Allocation Date', 'Expected Return Date', 'Actual Return Date', 'Purpose', 'Status']
    ws.append(headers)
    
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=4, column=col).font = header_font
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A4:J4"
    
    for obj in get_filtered_allocations(request):
        ws.append([
            obj.asset.name, obj.asset.asset_code, obj.department.name, obj.assigned_to, obj.quantity,
            obj.allocation_date.strftime("%Y-%m-%d"), 
            obj.expected_return_date.strftime("%Y-%m-%d") if obj.expected_return_date else '', 
            obj.actual_return_date.strftime("%Y-%m-%d") if obj.actual_return_date else '', 
            obj.purpose, obj.get_status_display()
        ])
        
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="allocations_report_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response

@login_required
@system_admin_required
def export_allocations_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="allocations_report_{datetime.now().strftime("%Y%m%d")}.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Campus Assets", styles['Heading1']))
    elements.append(Paragraph("Allocations Report", styles['Heading2']))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    data = [['Asset', 'Code', 'Department', 'Assigned To', 'Qty', 'Alloc Date', 'Exp Return', 'Act Return', 'Status']]
    for obj in get_filtered_allocations(request):
        data.append([
            Paragraph(obj.asset.name, styles['Normal']), obj.asset.asset_code, obj.department.name, obj.assigned_to or '-', str(obj.quantity),
            obj.allocation_date.strftime("%Y-%m-%d"), 
            obj.expected_return_date.strftime("%Y-%m-%d") if obj.expected_return_date else '-', 
            obj.actual_return_date.strftime("%Y-%m-%d") if obj.actual_return_date else '-', 
            obj.get_status_display()
        ])
        
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A3626')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    elements.append(t)
    doc.build(elements)
    return response

# -- Maintenance Exports --
@login_required
@system_admin_required
def export_maintenance_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="maintenance_report_{datetime.now().strftime("%Y%m%d")}.csv"'
    response.write('\ufeff'.encode('utf8'))
    
    writer = csv.writer(response)
    writer.writerow(['Asset', 'Asset Code', 'Description', 'Maintenance Date', 'Technician', 'Cost', 'Status'])
    
    for obj in get_filtered_maintenance(request):
        writer.writerow([
            obj.asset.name, obj.asset.asset_code, obj.description, obj.maintenance_date, obj.technician, obj.cost or '', obj.get_status_display()
        ])
    return response

@login_required
@system_admin_required
def export_maintenance_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Maintenance"
    
    ws.append(['Campus Assets - Maintenance Report'])
    ws.append([f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'])
    ws.append([])
    
    headers = ['Asset', 'Asset Code', 'Description', 'Maintenance Date', 'Technician', 'Cost', 'Status']
    ws.append(headers)
    
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=4, column=col).font = header_font
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A4:G4"
    
    for obj in get_filtered_maintenance(request):
        ws.append([
            obj.asset.name, obj.asset.asset_code, obj.description, 
            obj.maintenance_date.strftime("%Y-%m-%d"), obj.technician, obj.cost or 0.0, obj.get_status_display()
        ])
        ws.cell(row=ws.max_row, column=6).number_format = '[$₹-en-IN]#,##0.00'
        
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="maintenance_report_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response

@login_required
@system_admin_required
def export_maintenance_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="maintenance_report_{datetime.now().strftime("%Y%m%d")}.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Campus Assets", styles['Heading1']))
    elements.append(Paragraph("Maintenance Report", styles['Heading2']))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    data = [['Asset', 'Code', 'Description', 'Date', 'Technician', 'Cost', 'Status']]
    for obj in get_filtered_maintenance(request):
        cost_str = f"Rs. {obj.cost:.2f}" if obj.cost else "-"
        data.append([
            Paragraph(obj.asset.name, styles['Normal']), obj.asset.asset_code, 
            Paragraph(obj.description, styles['Normal']), 
            obj.maintenance_date.strftime("%Y-%m-%d"), obj.technician or '-', cost_str, obj.get_status_display()
        ])
        
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A3626')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    elements.append(t)
    doc.build(elements)
    return response

def home(request):
    if request.user.is_authenticated:
        return redirect('assets:dashboard')
    return render(request, 'assets/home.html')

def about(request):
    return render(request, 'assets/about.html')

def help_page(request):
    return render(request, 'assets/help.html')

def custom_404(request, exception=None):
    return render(request, '404.html', status=404)

def custom_403(request, exception=None):
    return render(request, '403.html', status=403)

def register(request):
    if request.user.is_authenticated:
        return redirect('assets:dashboard')
        
    if request.method == 'POST':
        form = UserRegistrationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.is_staff = False
            user.is_superuser = False
            user.save()
            profile, created = Profile.objects.get_or_create(user=user)
            profile.role = 'NORMAL_USER'
            profile.account_status = 'APPROVED'
            profile.save()
            messages.success(request, 'Your account has been created successfully. You can now sign in.')
            return redirect('assets:register_success')
        else:
            messages.error(request, 'Please correct the errors below.')
    else:
        form = UserRegistrationForm()
        
    return render(request, 'assets/register.html', {'form': form})

from django.contrib.auth.decorators import user_passes_test
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_POST

def register_success(request):
    return render(request, 'assets/registration_success.html')

@user_passes_test(lambda u: u.is_superuser)
@system_admin_required
def admin_management(request):
    profiles = Profile.objects.exclude(user=request.user).select_related('user')
    pending = profiles.filter(account_status='PENDING').order_by('-user__date_joined')
    approved = profiles.filter(account_status='APPROVED').order_by('-user__date_joined')
    rejected = profiles.filter(account_status='REJECTED').order_by('-user__date_joined')
    disabled = profiles.filter(account_status='DISABLED').order_by('-user__date_joined')
    
    context = {
        'pending_profiles': pending,
        'approved_profiles': approved,
        'rejected_profiles': rejected,
        'disabled_profiles': disabled,
        'total_admins': profiles.count() + 1, # +1 for current superuser
    }
    return render(request, 'assets/admin_management.html', context)

@require_POST
@user_passes_test(lambda u: u.is_superuser)
@system_admin_required
def admin_management_action(request, action, user_id):
    if request.user.id == user_id:
        messages.error(request, 'You cannot perform this action on your own account.')
        return redirect('assets:admin_management')
        
    profile = get_object_or_404(Profile, user__id=user_id)
    
    if action == 'approve':
        profile.account_status = 'APPROVED'
        profile.save()
        messages.success(request, 'Administrator approved successfully.')
    elif action == 'reject':
        profile.account_status = 'REJECTED'
        profile.save()
        messages.success(request, 'Administrator registration request rejected.')
    elif action == 'disable':
        profile.account_status = 'DISABLED'
        profile.save()
        messages.success(request, 'Administrator account disabled successfully.')
    elif action == 'enable':
        profile.account_status = 'APPROVED'
        profile.save()
        messages.success(request, 'Administrator account enabled successfully.')
    else:
        messages.error(request, 'Invalid action.')
        
    return redirect('assets:admin_management')


# ==========================================
# PHASE 3: DAMAGED EQUIPMENT VIEWS
# ==========================================

@login_required
def damaged_equipment_list(request):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    if role == 'NORMAL_USER':
        messages.error(request, "Access denied.")
        return redirect('assets:dashboard')
        
    if role == 'SYSTEM_ADMIN':
        damage_reports = DamagedEquipment.objects.select_related('asset', 'reported_by', 'assigned_technician').order_by('-created_at')
    else:
        # TECHNICIAN
        damage_reports = DamagedEquipment.objects.select_related('asset', 'reported_by', 'assigned_technician').order_by('-created_at')
        
    return render(request, 'assets/damaged_equipment_list.html', {
        'damage_reports': damage_reports,
        'severity_choices': DamagedEquipment.SEVERITY_CHOICES,
        'status_choices': DamagedEquipment.STATUS_CHOICES,
    })

@login_required
def damaged_equipment_update(request, pk):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    if role not in ['SYSTEM_ADMIN', 'TECHNICIAN']:
        messages.error(request, "Access denied.")
        return redirect('assets:dashboard')
        
    report = get_object_or_404(DamagedEquipment, pk=pk)
    
    if request.method == 'POST':
        status = request.POST.get('status')
        resolution = request.POST.get('resolution', '')
        
        # Technician validation
        if role == 'TECHNICIAN':
            # Technicians can only modify if they are assigned, or if unassigned
            if report.assigned_technician and report.assigned_technician != request.user:
                messages.error(request, "This report is assigned to another technician.")
                return redirect('assets:damaged_equipment_list')
            if not report.assigned_technician:
                report.assigned_technician = request.user
                
        if status:
            report.status = status
            
        report.resolution = resolution
        
        if status == 'RESOLVED' and not report.resolved_date:
            report.resolved_date = timezone.now().date()
            
        report.save()
        messages.success(request, "Damage report updated successfully.")
        
    return redirect('assets:damaged_equipment_list')

# ==========================================
# PHASE 4: USER REQUEST VIEWS
# ==========================================

@login_required
def user_requests(request):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    
    if role == 'NORMAL_USER':
        requests_list = UserRequest.objects.filter(user=request.user).select_related('asset', 'assigned_technician').order_by('-created_at')
        my_allocations = Allocation.objects.filter(assigned_user=request.user, status__in=['Active', 'Partially Returned']).select_related('asset')
        
        if request.method == 'POST':
            request_type = request.POST.get('request_type')
            subject = request.POST.get('subject')
            description = request.POST.get('description')
            asset_id = request.POST.get('asset_id')
            
            asset = None
            if request_type in ['Damage Report', 'Maintenance Request']:
                if not asset_id:
                    messages.error(request, f"Please select an asset for your {request_type}.")
                    return redirect('assets:user_requests')
                    
                # Security: Verify the asset is actively assigned to the user
                is_valid_asset = my_allocations.filter(asset_id=asset_id).exists()
                if not is_valid_asset:
                    messages.error(request, "Invalid asset selection. You can only report issues for equipment assigned to you.")
                    return redirect('assets:user_requests')
                asset = get_object_or_404(Asset, pk=asset_id)
            else:
                if asset_id:
                    # Allow referencing an assigned asset if they want, but optional
                    is_valid_asset = my_allocations.filter(asset_id=asset_id).exists()
                    if is_valid_asset:
                        asset = get_object_or_404(Asset, pk=asset_id)
            
            if request_type == 'Damage Report':
                # Create DamagedEquipment directly
                severity = request.POST.get('severity', 'LOW')
                DamagedEquipment.objects.create(
                    asset=asset,
                    reported_by=request.user,
                    damage_date=timezone.now().date(),
                    damage_description=description,
                    severity=severity,
                    status='REPORTED'
                )
                messages.success(request, "Damage report submitted successfully.")
            else:
                UserRequest.objects.create(
                    user=request.user,
                    asset=asset,
                    request_type=request_type,
                    subject=subject,
                    description=description,
                    status='PENDING'
                )
                messages.success(request, "Support request submitted successfully.")
                
            return redirect('assets:user_requests')
            
        return render(request, 'assets/user_requests.html', {
            'requests_list': requests_list,
            'my_allocations': my_allocations,
            'request_type_choices': UserRequest.REQUEST_TYPE_CHOICES,
            'severity_choices': DamagedEquipment.SEVERITY_CHOICES
        })
        
    elif role == 'TECHNICIAN':
        return redirect('assets:technician_requests')
    else:
        # SYSTEM ADMIN
        requests_list = UserRequest.objects.select_related('user', 'asset', 'assigned_technician').order_by('-created_at')
        return render(request, 'assets/admin_requests.html', {
            'requests_list': requests_list,
            'status_choices': UserRequest.STATUS_CHOICES
        })

@login_required
def technician_requests(request):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    if role not in ['TECHNICIAN', 'SYSTEM_ADMIN']:
        messages.error(request, "Access denied.")
        return redirect('assets:dashboard')
        
    requests_list = UserRequest.objects.select_related('user', 'asset', 'assigned_technician').order_by('-created_at')
    
    return render(request, 'assets/technician_requests.html', {
        'requests_list': requests_list,
        'status_choices': UserRequest.STATUS_CHOICES
    })

@login_required
def request_update(request, pk):
    role = getattr(request.user, 'profile', None) and request.user.profile.role
    if role not in ['SYSTEM_ADMIN', 'TECHNICIAN']:
        messages.error(request, "Access denied.")
        return redirect('assets:dashboard')
        
    req = get_object_or_404(UserRequest, pk=pk)
    
    if request.method == 'POST':
        status = request.POST.get('status')
        technician_response = request.POST.get('technician_response', '')
        
        # Technician validation
        if role == 'TECHNICIAN':
            if req.assigned_technician and req.assigned_technician != request.user:
                messages.error(request, "This request is assigned to another technician.")
                return redirect('assets:technician_requests')
            if not req.assigned_technician:
                req.assigned_technician = request.user
                
        if status:
            req.status = status
            
        if technician_response:
            req.technician_response = technician_response
            
        if status == 'RESOLVED' and not req.resolved_at:
            req.resolved_at = timezone.now()
            
        req.save()
        messages.success(request, "Request updated successfully.")
        
    if role == 'SYSTEM_ADMIN':
        return redirect('assets:user_requests') # admin uses user_requests URL as its base
    return redirect('assets:technician_requests')

# ==========================================
# PHASE 5: FEEDBACK VIEWS
# ==========================================

@login_required
def submit_feedback(request, pk):
    req = get_object_or_404(UserRequest, pk=pk)
    
    # Validation: request belongs to logged-in user
    if req.user != request.user:
        messages.error(request, "You can only submit feedback for your own requests.")
        return redirect('assets:user_requests')
        
    # Validation: request is resolved
    if req.status != 'RESOLVED':
        messages.error(request, "You can only submit feedback for resolved requests.")
        return redirect('assets:user_requests')
        
    # Validation: feedback does not already exist
    if hasattr(req, 'feedback'):
        messages.error(request, "Feedback has already been submitted for this request.")
        return redirect('assets:user_requests')
        
    if request.method == 'POST':
        rating = request.POST.get('rating')
        feedback_text = request.POST.get('feedback', '')
        
        if not rating or not str(rating).isdigit() or not (1 <= int(rating) <= 5):
            messages.error(request, "Invalid rating value.")
            return redirect('assets:user_requests')
            
        UserFeedback.objects.create(
            request=req,
            user=request.user,
            rating=int(rating),
            feedback=feedback_text
        )
        messages.success(request, "Thank you for your feedback!")
        
    return redirect('assets:user_requests')

# ==========================================
# ADMIN PEOPLE & FEEDBACK MANAGEMENT
# ==========================================

@login_required
@system_admin_required
def user_list(request):
    users = User.objects.filter(profile__role='NORMAL_USER').select_related('profile').order_by('-date_joined')
    # Count assigned equipment, requests
    users = users.annotate(
        assigned_equipment_count=Count('allocations', filter=Q(allocations__status__in=['Active', 'Partially Returned'])),
        request_count=Count('support_requests')
    )
    return render(request, 'assets/user_list.html', {'users': users})

@login_required
@system_admin_required
def technician_list(request):
    technicians = User.objects.filter(profile__role='TECHNICIAN').select_related('profile').order_by('-date_joined')
    # Count requests, damage reports assigned
    technicians = technicians.annotate(
        assigned_requests_count=Count('assigned_support_requests'),
        assigned_damage_count=Count('assigned_damage_reports')
    )
    return render(request, 'assets/technician_list.html', {'technicians': technicians})

@require_POST
@login_required
@system_admin_required
def user_action(request, action, user_id):
    target_user = get_object_or_404(User, pk=user_id)
    
    # Security: Only operate on NORMAL_USER
    if not hasattr(target_user, 'profile') or target_user.profile.role != 'NORMAL_USER':
        messages.error(request, "Invalid operation. Target is not a Normal User.")
        return redirect('assets:user_list')
        
    if action == 'enable':
        target_user.is_active = True
        target_user.profile.account_status = 'APPROVED'
        messages.success(request, f"User {target_user.username} enabled.")
    elif action == 'disable':
        target_user.is_active = False
        target_user.profile.account_status = 'REJECTED'
        messages.success(request, f"User {target_user.username} disabled.")
        
    target_user.save()
    target_user.profile.save()
    
    return redirect('assets:user_list')

@require_POST
@login_required
@system_admin_required
def technician_action(request, action, user_id):
    target_user = get_object_or_404(User, pk=user_id)
    
    # Security: Do not allow disabling self
    if target_user == request.user:
        messages.error(request, "You cannot disable your own account.")
        return redirect('assets:technician_list')
        
    # Security: Only operate on TECHNICIAN
    if not hasattr(target_user, 'profile') or target_user.profile.role != 'TECHNICIAN':
        messages.error(request, "Invalid operation. Target is not a Technician.")
        return redirect('assets:technician_list')
        
    if action == 'enable':
        target_user.is_active = True
        target_user.profile.account_status = 'APPROVED'
        messages.success(request, f"Technician {target_user.username} enabled.")
    elif action == 'disable':
        target_user.is_active = False
        target_user.profile.account_status = 'REJECTED'
        messages.success(request, f"Technician {target_user.username} disabled.")
        
    target_user.save()
    target_user.profile.save()
    
    return redirect('assets:technician_list')

@login_required
@system_admin_required
def feedback_list(request):
    feedbacks = UserFeedback.objects.select_related('request', 'user', 'request__assigned_technician', 'request__asset').order_by('-created_at')
    
    total_feedback = feedbacks.count()
    avg_rating = feedbacks.aggregate(avg=Avg('rating'))['avg'] or 0
    
    context = {
        'feedbacks': feedbacks,
        'total_feedback': total_feedback,
        'avg_rating': round(avg_rating, 1)
    }
    return render(request, 'assets/feedback_list.html', context)
